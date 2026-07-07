"""Excel design-sheet adapter.

The design studio already produces an Excel per order — this adapter taps
that workflow instead of changing it. Point a folder watcher at the studio's
shared folder and every saved sheet lands in the database within minutes.

Template: key/value pairs in columns A/B of the first worksheet. Labels are
matched case-insensitively; unknown labels are ignored; missing values stay
None. Like every ingestion door, output is the validated pydantic contract —
downstream code never sees a spreadsheet.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..schema import Article, Construction, WeftSpin

# spreadsheet label (lowercased) -> Construction field
_FIELD_MAP = {
    "warp count": "warp_count",
    "weft count (ne)": "weft_count_ne",
    "weft denier": "weft_denier",
    "weft spin": "weft_spin",
    "weft material": "weft_material",
    "epi": "epi",
    "ppi": "ppi",
    "reed count": "reed_count",
    "pile ratio": "pile_ratio",
    "gsm": "gsm",
    "width (cm)": "width_cm",
    "selvedge type": "selvedge_type",
}


def parse_design_sheet(path: Path) -> Article:
    """Read one design-sheet Excel into a validated Article."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]

    raw: dict[str, object] = {}
    article_id: str | None = None
    for row in sheet.iter_rows(min_col=1, max_col=2, values_only=True):
        label, value = row[0], row[1]
        if label is None or value is None:
            continue
        key = str(label).strip().lower()
        if key == "article id":
            article_id = str(value).strip()
            continue
        field = _FIELD_MAP.get(key)
        if field is not None:
            raw[field] = value
    workbook.close()

    if not article_id:
        raise ValueError(f"{path.name}: 'Article ID' row is required")

    if "weft_spin" in raw:
        spin = str(raw["weft_spin"]).strip().lower()
        aliases = {"open end": "oe", "open-end": "oe", "rotor": "oe"}
        raw["weft_spin"] = WeftSpin(aliases.get(spin, spin))
    if "weft_material" in raw:
        raw["weft_material"] = str(raw["weft_material"]).strip().lower()

    return Article(article_id=article_id, construction=Construction.model_validate(raw))
